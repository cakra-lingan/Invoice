import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

class InvoiceApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Aplikasi Invoice - Irawan")
        self.root.geometry("600x700")
        self.root.configure(bg="#f0f0f0")
        
        self.items = []  # List untuk menyimpan item barang
        self.excel_file = "invoices.xlsx"
        self.init_excel()
        
        # ===== HEADER =====
        header = tk.Label(root, text="📋 SISTEM INVOICE IRAWAN", 
                         font=("Arial", 16, "bold"), bg="#2c3e50", fg="white", pady=10)
        header.pack(fill=tk.X)
        
        # ===== FRAME UTAMA =====
        main_frame = ttk.Frame(root, padding=20)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # --- Bagian 1: Data Pembeli ---
        ttk.Label(main_frame, text="DATA PEMBELI", font=("Arial", 12, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0,10))
        
        ttk.Label(main_frame, text="Nama Pembeli:").grid(row=1, column=0, sticky="w", pady=5)
        self.pembeli_entry = ttk.Entry(main_frame, width=40)
        self.pembeli_entry.grid(row=1, column=1, sticky="ew", pady=5)
        
        ttk.Label(main_frame, text="Alamat:").grid(row=2, column=0, sticky="w", pady=5)
        self.alamat_entry = ttk.Entry(main_frame, width=40)
        self.alamat_entry.grid(row=2, column=1, sticky="ew", pady=5)
        
        ttk.Label(main_frame, text="No. Telepon:").grid(row=3, column=0, sticky="w", pady=5)
        self.telepon_entry = ttk.Entry(main_frame, width=40)
        self.telepon_entry.grid(row=3, column=1, sticky="ew", pady=5)
        
        # --- Bagian 2: Barang/Jasa ---
        ttk.Label(main_frame, text="DETAIL BARANG/JASA", font=("Arial", 12, "bold")).grid(row=4, column=0, columnspan=2, sticky="w", pady=(15,10))
        
        ttk.Label(main_frame, text="Deskripsi:").grid(row=5, column=0, sticky="w", pady=5)
        self.deskripsi_entry = ttk.Entry(main_frame, width=40)
        self.deskripsi_entry.grid(row=5, column=1, sticky="ew", pady=5)
        
        ttk.Label(main_frame, text="Harga Satuan (Rp):").grid(row=6, column=0, sticky="w", pady=5)
        self.harga_entry = ttk.Entry(main_frame, width=40)
        self.harga_entry.grid(row=6, column=1, sticky="ew", pady=5)
        
        ttk.Label(main_frame, text="Jumlah (Qty):").grid(row=7, column=0, sticky="w", pady=5)
        self.qty_entry = ttk.Entry(main_frame, width=40)
        self.qty_entry.grid(row=7, column=1, sticky="ew", pady=5)
        
        # Tombol Tambah Item
        ttk.Button(main_frame, text="➕ Tambah Item", command=self.tambah_item).grid(row=8, column=0, columnspan=2, sticky="ew", pady=10)
        
        # --- Daftar Item ---
        ttk.Label(main_frame, text="ITEM YANG DITAMBAHKAN:", font=("Arial", 10, "bold")).grid(row=9, column=0, columnspan=2, sticky="w", pady=(10,5))
        
        self.item_listbox = tk.Listbox(main_frame, height=6, width=60, bg="white", border=1)
        self.item_listbox.grid(row=10, column=0, columnspan=2, sticky="ew", pady=5)
        
        # Scrollbar untuk listbox
        scrollbar = ttk.Scrollbar(main_frame, orient=tk.VERTICAL, command=self.item_listbox.config)
        scrollbar.grid(row=10, column=2, sticky="ns", pady=5)
        self.item_listbox.config(yscrollcommand=scrollbar.set)
        
        # Tombol Hapus Item
        ttk.Button(main_frame, text="🗑️ Hapus Item Terpilih", command=self.hapus_item).grid(row=11, column=0, columnspan=2, sticky="ew", pady=5)
        
        # --- Total ---
        ttk.Label(main_frame, text="TOTAL INVOICE (Rp):", font=("Arial", 11, "bold")).grid(row=12, column=0, sticky="w", pady=(10,5))
        self.total_label = ttk.Label(main_frame, text="0", font=("Arial", 11, "bold"), foreground="green")
        self.total_label.grid(row=12, column=1, sticky="w", pady=(10,5))
        
        # --- Tombol Aksi ---
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=13, column=0, columnspan=2, sticky="ew", pady=(20,0))
        
        ttk.Button(button_frame, text="💾 BUAT INVOICE", command=self.buat_invoice).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="📂 Buka File Excel", command=self.buka_excel).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="🔄 Reset", command=self.reset_form).pack(side=tk.LEFT, padx=5)
        
        main_frame.columnconfigure(1, weight=1)
    
    def init_excel(self):
        """Inisialisasi file Excel jika belum ada"""
        if not os.path.exists(self.excel_file):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Invoices"
            
            # Header
            headers = ["No. Invoice", "Tanggal", "Pembeli", "Alamat", "Telepon", "Deskripsi", "Harga Satuan", "Qty", "Subtotal", "Total"]
            ws.append(headers)
            
            # Format header
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            wb.save(self.excel_file)
    
    def tambah_item(self):
        """Tambah item ke list"""
        deskripsi = self.deskripsi_entry.get().strip()
        harga = self.harga_entry.get().strip()
        qty = self.qty_entry.get().strip()
        
        if not deskripsi or not harga or not qty:
            messagebox.showwarning("Input Tidak Lengkap", "Mohon isi semua field barang/jasa!")
            return
        
        try:
            harga = int(harga)
            qty = int(qty)
            subtotal = harga * qty
        except ValueError:
            messagebox.showerror("Input Salah", "Harga dan Qty harus angka!")
            return
        
        self.items.append({
            'deskripsi': deskripsi,
            'harga': harga,
            'qty': qty,
            'subtotal': subtotal
        })
        
        # Tampilkan di listbox
        display_text = f"{deskripsi} | Rp{harga:,} x {qty} = Rp{subtotal:,}"
        self.item_listbox.insert(tk.END, display_text)
        
        # Update total
        self.update_total()
        
        # Clear input
        self.deskripsi_entry.delete(0, tk.END)
        self.harga_entry.delete(0, tk.END)
        self.qty_entry.delete(0, tk.END)
        self.deskripsi_entry.focus()
    
    def hapus_item(self):
        """Hapus item yang dipilih"""
        selection = self.item_listbox.curselection()
        if not selection:
            messagebox.showwarning("Tidak Ada Pilihan", "Pilih item yang ingin dihapus!")
            return
        
        idx = selection[0]
        self.item_listbox.delete(idx)
        del self.items[idx]
        self.update_total()
    
    def update_total(self):
        """Update total invoice"""
        total = sum(item['subtotal'] for item in self.items)
        self.total_label.config(text=f"{total:,}")
    
    def generate_invoice_number(self):
        """Generate nomor invoice acak INV-XXXX"""
        return f"INV-{random.randint(1000, 9999)}"
    
    def buat_invoice(self):
        """Buat dan simpan invoice ke Excel"""
        pembeli = self.pembeli_entry.get().strip()
        alamat = self.alamat_entry.get().strip()
        telepon = self.telepon_entry.get().strip()
        
        if not pembeli or not alamat or not telepon:
            messagebox.showwarning("Data Tidak Lengkap", "Mohon isi semua data pembeli!")
            return
        
        if not self.items:
            messagebox.showwarning("Tidak Ada Item", "Mohon tambah minimal 1 item barang/jasa!")
            return
        
        try:
            invoice_number = self.generate_invoice_number()
            tanggal = datetime.now().strftime("%d/%m/%Y")
            total = sum(item['subtotal'] for item in self.items)
            
            wb = openpyxl.load_workbook(self.excel_file)
            ws = wb.active
            
            # Simpan data untuk setiap item
            for idx, item in enumerate(self.items):
                row_data = [
                    invoice_number if idx == 0 else "",
                    tanggal if idx == 0 else "",
                    pembeli if idx == 0 else "",
                    alamat if idx == 0 else "",
                    telepon if idx == 0 else "",
                    item['deskripsi'],
                    item['harga'],
                    item['qty'],
                    item['subtotal'],
                    total if idx == 0 else ""
                ]
                ws.append(row_data)
            
            # Format cells
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=10):
                for cell in row:
                    cell.border = border
                    if cell.column in [7, 8, 9, 10]:  # Kolom angka
                        cell.number_format = '#,##0'
            
            # Set lebar kolom
            ws.column_dimensions['A'].width = 15
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 20
            ws.column_dimensions['E'].width = 15
            ws.column_dimensions['F'].width = 20
            ws.column_dimensions['G'].width = 15
            ws.column_dimensions['H'].width = 8
            ws.column_dimensions['I'].width = 15
            ws.column_dimensions['J'].width = 15
            
            wb.save(self.excel_file)
            
            messagebox.showinfo("Sukses!", f"Invoice {invoice_number} berhasil dibuat!\nData tersimpan di {self.excel_file}")
            self.reset_form()
            
        except Exception as e:
            messagebox.showerror("Error", f"Gagal membuat invoice: {str(e)}")
    
    def buka_excel(self):
        """Buka file Excel"""
        if os.path.exists(self.excel_file):
            os.startfile(self.excel_file)
        else:
            messagebox.showinfo("Info", "File invoice belum ada. Buat invoice terlebih dahulu!")
    
    def reset_form(self):
        """Reset semua form"""
        self.pembeli_entry.delete(0, tk.END)
        self.alamat_entry.delete(0, tk.END)
        self.telepon_entry.delete(0, tk.END)
        self.deskripsi_entry.delete(0, tk.END)
        self.harga_entry.delete(0, tk.END)
        self.qty_entry.delete(0, tk.END)
        self.item_listbox.delete(0, tk.END)
        self.items = []
        self.update_total()
        self.pembeli_entry.focus()

if __name__ == "__main__":
    root = tk.Tk()
    app = InvoiceApp(root)
    root.mainloop()